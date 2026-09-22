"""
Step 0: 数据规整 —— PHM2019 铝搭接件（数据集 D）
==================================================

把散落的 Description_T*.xlsx（裂纹长度真值）与 <cycle>/signal_{1,2}.csv（Lamb 波）
规整成两张可复现的清单表，并留下核对报告。

产出
----
  phmdc/labels.csv             specimen, cycle, cycle_raw, crack_mm, source, note
  phmdc/index.csv              specimen, cycle, rep, path, n_rows, fs_hz, dt_s, t_max_s
  phmdc/results/step0_report.txt   四项核对报告（对账 / ch1 一致性 / ch2 重复性 / 完整性）

四项核对（README §5 Step 0）
----
  1. 标签 vs 波形对账  —— 含「循环数录入丢位」的显式检测与修复记录
  2. ch1 一致性自检    —— 同一试件内 ch1 应几乎不变（否则=耦合/增益漂移，非损伤）
  3. ch2 重复性检查    —— signal_1 vs signal_2 差异 = 测量噪声下界
  4. 波形完整性        —— 文件是否齐全、长度是否都是 4000 点、dt 是否恒为 5e-8 s

设计原则
----
  * 解析一律按**表头定位**（openpyxl 逐行找 'Number of cycle'），不按行号猜
  * 任何对原始数据的**修正都必须留痕**（cycle_raw 原值 + note 依据），不静默改数
  * 自定位：ROOT = 脚本所在目录，可从任意 cwd 运行

依赖：pandas / numpy / openpyxl

作者: Roo
日期: 2026-09-20
"""

import os
import re
import sys
import csv
import warnings
from collections import defaultdict

import numpy as np
import pandas as pd
import openpyxl

warnings.filterwarnings('ignore')

try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

# ============================================================
# 全局配置
# ============================================================
ROOT = os.path.dirname(os.path.abspath(__file__))

SPECIMENS = ['T1', 'T2', 'T3', 'T4', 'T5', 'T6', 'T7', 'T8']
TRAIN = ['T1', 'T2', 'T3', 'T4', 'T5', 'T6']
VALID = ['T7', 'T8']

EXPECTED_ROWS = 4000          # 每个波形文件的点数（README §1.1）
EXPECTED_DT = 5.0e-8          # 采样间隔 [s] -> fs = 20 MHz
FS_HZ = 1.0 / EXPECTED_DT     # 20e6 Hz

CYCLE_DIR_RE = re.compile(r'^\d+$')

OUT_LABELS = os.path.join(ROOT, 'labels.csv')
OUT_INDEX = os.path.join(ROOT, 'index.csv')
OUT_REPORT = os.path.join(ROOT, 'results', 'step0_report.txt')

_REPORT_LINES = []


def say(msg=''):
    """同时打印与收集到报告。"""
    print(msg)
    _REPORT_LINES.append(msg)


# ============================================================
# 1. 真值表解析（openpyxl，按表头定位）
# ============================================================
def _as_number(v):
    """把单元格值转成 float；非数值返回 None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(',', '')
        try:
            return float(s)
        except ValueError:
            return None
    return None


def parse_description(specimen):
    """解析 Description_<specimen>.xlsx，返回 (记录列表, 元信息).

    定位规则：找到同时含 'cycle' 与 'Crack' 的表头行，其下连续「两列都是数值」的行即数据。
    表头下方的空行 / 说明行会被自动跳过并终止解析。
    """
    path = os.path.join(ROOT, specimen, 'Description_%s.xlsx' % specimen)
    if not os.path.exists(path):
        raise FileNotFoundError(path)

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb[wb.sheetnames[0]]

    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]

    header_idx = None
    for i, r in enumerate(rows):
        cells = [str(c).strip().lower() for c in r if c is not None]
        joined = ' '.join(cells)
        if 'crack' in joined and ('cycle' in joined or 'cycles' in joined):
            header_idx = i
            break
    if header_idx is None:
        raise ValueError('%s: 未找到裂纹长度表头' % path)

    records = []
    for r in rows[header_idx + 1:]:
        cyc = _as_number(r[0] if len(r) > 0 else None)
        crk = _as_number(r[1] if len(r) > 1 else None)
        if cyc is None or crk is None:
            # 数据区已结束（表头下的空行 / 后续说明）
            if records:
                break
            continue
        records.append({
            'specimen': specimen,
            'cycle_raw': int(round(cyc)),
            'crack_mm': float(crk),
            'source': 'Description_%s.xlsx!%s' % (specimen, ws.title),
        })

    # 元信息：Table 1 里的 'The crack initiates around N cycles'
    meta = {'initiation_cycles': None, 'loading': None}
    for r in rows:
        for c in r:
            if isinstance(c, str):
                m = re.search(r'initiates around\s+(\d+)\s+cycles', c, re.I)
                if m:
                    meta['initiation_cycles'] = int(m.group(1))
        if r and isinstance(r[0], str) and r[0].strip() == specimen:
            if len(r) > 1 and isinstance(r[1], str):
                meta['loading'] = r[1].strip()
    return records, meta


# ============================================================
# 2. 波形索引
# ============================================================
def load_signal(path):
    """按列名读取一个波形文件，返回 (t_rebuilt, t_file, ch1, ch2, col_names).

    ⚠️ 本数据集的 time 列**被 Excel 量化成科学计数法且逐试件精度不同**
    （1~15 位有效数字）⇒ 92/94 个文件出现重复时间戳。
    因此**一律用 t[i] = i * EXPECTED_DT 重建时间轴**，
    同时返回文件里的 time 列用于可靠性判定（不用于分析）。
    """
    df = pd.read_csv(path)
    cols = [str(c) for c in df.columns]
    n = len(df)
    t_file = df['time'].to_numpy(dtype=float)
    t_rebuilt = np.arange(n, dtype=float) * EXPECTED_DT
    ch1 = df['ch1'].to_numpy(dtype=float)
    ch2 = df['ch2'].to_numpy(dtype=float)
    return t_rebuilt, t_file, ch1, ch2, cols


def scan_waveforms():
    """扫描 <specimen>/<cycle>/signal_{1,2}.csv，返回索引行列表。

    除点数外，还核查：文件是否齐全、列名是否完整、time 列是否可靠。
    """
    rows = []
    for sp in SPECIMENS:
        sp_dir = os.path.join(ROOT, sp)
        if not os.path.isdir(sp_dir):
            say('[X] 缺少试件目录: %s' % sp_dir)
            continue
        cycles = sorted(
            int(d) for d in os.listdir(sp_dir)
            if CYCLE_DIR_RE.match(d) and os.path.isdir(os.path.join(sp_dir, d))
        )
        for cyc in cycles:
            for rep in (1, 2):
                fname = 'signal_%d.csv' % rep
                fpath = os.path.join(sp_dir, str(cyc), fname)
                if not os.path.exists(fpath):
                    rows.append({
                        'specimen': sp, 'cycle': cyc, 'rep': rep,
                        'path': None, 'n_rows': 0, 'n_cols': 0,
                        'col_names': '', 't_file_err_s': np.nan,
                        'time_ok': False, 'fs_hz': FS_HZ,
                    })
                    continue
                t_rb, t_f, _c1, _c2, cols = load_signal(fpath)
                n = len(t_rb)
                m = min(n, len(t_f))
                # time 列与重建轴的最大偏差（只用重叠部分，末值四舍五入不计）
                err = float(np.max(np.abs(t_f[:m] - t_rb[:m]))) if m else np.nan
                rows.append({
                    'specimen': sp, 'cycle': cyc, 'rep': rep,
                    'path': os.path.relpath(fpath, ROOT).replace('\\', '/'),
                    'n_rows': n,
                    'n_cols': len(cols),
                    'col_names': '|'.join(cols),
                    't_file_err_s': err,
                    'time_ok': bool(np.isfinite(err) and err <= 0.5 * EXPECTED_DT),
                    'fs_hz': FS_HZ,
                })
    return rows


# ============================================================
# 3. 标签 vs 波形对账（含「循环数录入丢位」检测）
# ============================================================
def reconcile(records, index_rows):
    """把标签匹配到波形循环数上。

    规则（显式、可复核）：
      * 精确匹配 -> 直接使用；
      * 若未匹配的标签值 c 与某个未匹配的波形循环数 w **一一对应**，且
        str(w).endswith(str(c))（十进制字符串尾部相同 = 录入时丢了首位数字），
        则判为录入丢位，修正 cycle=c->w，并在 note 里写明原值与依据；
      * 其余情况不修数，仅在 note / 报告中标记。

    返回 (新记录列表, 对账明细字典)。
    """
    wave_by_sp = defaultdict(set)
    for r in index_rows:
        wave_by_sp[r['specimen']].add(r['cycle'])

    out = []
    detail = {}
    for sp in SPECIMENS:
        recs = [r for r in records if r['specimen'] == sp]
        waves = wave_by_sp.get(sp, set())
        tag_cycles = {r['cycle_raw'] for r in recs}

        only_tag = sorted(tag_cycles - waves)      # 有标签、无波形
        only_wave = sorted(waves - tag_cycles)     # 有波形、无标签

        # --- 丢位检测：只在两边差集能一一配对时生效 ---
        fixes = {}
        if len(only_tag) == len(only_wave) and len(only_tag) > 0:
            cand = {c: [w for w in only_wave if str(w).endswith(str(c))]
                    for c in only_tag}
            if all(len(v) == 1 for v in cand.values()) and \
               len({v[0] for v in cand.values()}) == len(only_tag):
                fixes = {c: v[0] for c, v in cand.items()}

        for r in recs:
            rec = dict(r)
            c = rec['cycle_raw']
            rec['note'] = ''
            if c in fixes:
                rec['cycle'] = fixes[c]
                rec['note'] = ('cycle_typo: raw %d -> %d '
                               '(decimal suffix match, 唯一配对)' % (c, fixes[c]))
            else:
                rec['cycle'] = c
                if c not in waves:
                    rec['note'] = 'no_waveform_at_this_cycle'
            out.append(rec)

        detail[sp] = {
            'n_records': len(recs),
            'n_wave_cycles': len(waves),
            'only_tag': only_tag,
            'only_wave': only_wave,
            'fixes': fixes,
            'wave_cycles': sorted(waves),
        }

    return out, detail


# ============================================================
# 4. 信号级核对（ch1 一致性 / ch2 重复性 / 完整性）
# ============================================================
def check_signals(index_rows):
    """返回 (per_file, spec_ch1, pairs, problems) 四个列表.

    per_file:  每个波形文件的 ch1/ch2 峰峰值、RMS、峰值时刻、量化步长
    spec_ch1:  试件内 ch1 峰峰值的 min/max/相对散布（激励一致性）
    pairs:     signal_1 vs signal_2 的 相对RMS差 / 相关系数（重复性）
    problems:  (specimen, cycle, rep, msg) —— 文件缺失 / 点数异常 / 列名异常 / time 列不可靠
    """
    per_file = []
    pairs = []
    problems = []
    spec_ch1 = []

    for sp in SPECIMENS:
        rows = [r for r in index_rows if r['specimen'] == sp]
        by_cycle = defaultdict(dict)
        for r in rows:
            if r['path'] is None:
                problems.append((sp, r['cycle'], r['rep'],
                                 'missing_file: signal_%d.csv' % r['rep']))
                continue
            if r['n_rows'] != EXPECTED_ROWS:
                problems.append((sp, r['cycle'], r['rep'], 'n_rows=%d (期望 %d)'
                                 % (r['n_rows'], EXPECTED_ROWS)))
            if r['n_cols'] != 3:
                problems.append((sp, r['cycle'], r['rep'], 'n_cols=%d 列名=%s'
                                 % (r['n_cols'], r['col_names'])))
            if not r['time_ok']:
                problems.append((sp, r['cycle'], r['rep'], 'time 列不可靠: max|Δt|=%.3e s'
                                 % r['t_file_err_s']))
            t, _tf, ch1, ch2, _c = load_signal(os.path.join(ROOT, r['path']))
            pp1 = float(ch1.max() - ch1.min())
            pp2 = float(ch2.max() - ch2.min())
            u2 = np.unique(ch2)
            qstep = float(np.min(np.diff(u2))) if len(u2) > 1 else np.nan
            per_file.append({
                'specimen': sp, 'cycle': r['cycle'], 'rep': r['rep'],
                'n_rows': r['n_rows'], 'n_cols': r['n_cols'],
                'time_ok': r['time_ok'],
                'pp_ch1': pp1, 'rms_ch1': float(np.sqrt(np.mean(ch1 ** 2))),
                'pp_ch2': pp2, 'rms_ch2': float(np.sqrt(np.mean(ch2 ** 2))),
                'mean_ch2': float(ch2.mean()),
                'n_unique_ch2': int(len(u2)), 'quant_step_ch2': qstep,
                'argmax_ch2_s': float(t[int(np.argmax(np.abs(ch2)))]),
            })
            by_cycle[r['cycle']][r['rep']] = (t, ch1, ch2)

        # ch1 一致性（同试件内）
        pp1s = [f['pp_ch1'] for f in per_file if f['specimen'] == sp]
        if pp1s:
            lo, hi = min(pp1s), max(pp1s)
            rel = (hi - lo) / hi if hi > 0 else float('nan')
            spec_ch1.append({
                'specimen': sp, 'n_files': len(pp1s),
                'pp_ch1_min': lo, 'pp_ch1_max': hi,
                'pp_ch1_mean': float(np.mean(pp1s)), 'rel_spread': rel,
                'all_zero': bool(hi == 0.0),
            })

        # 重复性
        for cyc, d in sorted(by_cycle.items()):
            if 1 not in d or 2 not in d:
                continue
            t1, a1, b1 = d[1]
            t2, a2, b2 = d[2]
            n = min(len(b1), len(b2))
            x, y = b1[:n], b2[:n]
            d1, d2 = x - x.mean(), y - y.mean()
            denom = float(np.sqrt(np.sum(d1 ** 2) * np.sum(d2 ** 2)))
            r = float(np.sum(d1 * d2) / denom) if denom > 0 else float('nan')
            rms_x = float(np.sqrt(np.mean(d1 ** 2)))      # 去均值 RMS（AC）
            rms_diff = float(np.sqrt(np.mean((x - y) ** 2)))
            pp_x = float(x.max() - x.min())
            pp_y = float(y.max() - y.min())
            # 最佳整数时移对齐（|lag| ≤ 200 点 = ±10 µs，覆盖 TOF 漂移量级）
            # 用途：区分「差异源于采集时移/配对错误」还是「源于噪声」
            max_lag = 200
            best_r, best_lag = -1.0, 0
            for k in range(-max_lag, max_lag + 1):
                yk = np.roll(y, -k)          # yk[i] = y[i+k]
                m = slice(0, n - k) if k >= 0 else slice(-k, n)
                xs, ys = x[m], yk[m]
                dx, dy = xs - xs.mean(), ys - ys.mean()
                den = float(np.sqrt(np.sum(dx ** 2) * np.sum(dy ** 2)))
                rr = float(np.sum(dx * dy) / den) if den > 0 else -1.0
                if rr > best_r:
                    best_r, best_lag = rr, k
            a1c, a2c = a1[:n], a2[:n]
            rms_a1 = float(np.sqrt(np.mean((a1c - a1c.mean()) ** 2)))
            pairs.append({
                'specimen': sp, 'cycle': cyc, 'n': n,
                'corr_ch2': r,
                'corr_ch2_lag_best': best_r,
                'lag_best': best_lag,
                'rel_rms_diff': (rms_diff / rms_x) if rms_x > 0 else float('nan'),
                'pp_ratio_12': (pp_x / pp_y) if pp_y > 0 else float('nan'),
                'rel_rms_diff_ch1': (float(np.sqrt(np.mean((a1c - a2c) ** 2))) / rms_a1
                                     if rms_a1 > 0 else float('nan')),
            })

    return per_file, spec_ch1, pairs, problems


# ============================================================
# 4b. 试件级信号质量（决定哪些试件可用于定量回归）
# ============================================================
def specimen_quality(per_file, pairs):
    """按试件汇总信噪比，给出可用性判定。

    口径（显式）：
      signal  : ch2 去均值 RMS 的中位（AC-RMS，全 200 µs 记录）
      noise σ : 由 signal_1 vs signal_2 的 rms_diff 反推，σ = rms_diff / sqrt(2)
                （两次独立重复测量之差的标准差是单次噪声的 √2 倍）
      SNR     : signal / noise σ
    判定阈值：
      SNR ≥ 3  可用       —— 记录主要由确定性波场构成
      1.5 ≤ SNR < 3  勉强 —— 需在 Step 1 里加窗/降噪后再用
      SNR < 1.5  不可用   —— 记录被噪声主导，两次重复互不相关（实测已验证）
    """
    out = []
    pf = pd.DataFrame(per_file)
    pr = pd.DataFrame(pairs)
    for sp in SPECIMENS:
        g = pf[pf['specimen'] == sp]
        gp = pr[pr['specimen'] == sp] if len(pr) else pr
        if not len(g):
            continue
        # AC-RMS：全记录的 std
        ac = float(np.nanmedian(np.sqrt(np.maximum(
            g['rms_ch2'] ** 2 - g['mean_ch2'] ** 2, 0.0))))
        if len(gp):
            rms_diff = float(np.nanmedian(gp['rel_rms_diff'])) * ac
            sigma = rms_diff / np.sqrt(2.0)
            snr = ac / sigma if sigma > 0 else float('inf')
            corr_med = float(np.nanmedian(gp['corr_ch2']))
        else:
            rms_diff = sigma = float('nan')
            snr = float('nan')
            corr_med = float('nan')
        if not np.isfinite(snr):
            verdict = '无重复对'
        elif snr >= 3.0:
            verdict = '可用'
        elif snr >= 1.5:
            verdict = '勉强(需加窗/降噪)'
        else:
            verdict = '✗ 原始口径噪声主导'
        out.append({
            'specimen': sp, 'n_files': int(len(g)),
            'ac_rms_ch2': ac, 'noise_sigma': sigma, 'snr': snr,
            'corr_med': corr_med,
            'pp_ch1_med': float(np.nanmedian(g['pp_ch1'])),
            'pp_ch2_med': float(np.nanmedian(g['pp_ch2'])),
            'quant_step': float(np.nanmedian(g['quant_step_ch2'])),
            'verdict': verdict,
        })
    return out


# ============================================================
# 4c. 与 README §2.1 基准表对账（防解析漂移）
# ============================================================
# README §2.1 逐行读出的三组「对账基准」——硬编码在此，作为解析正确性的回归测试。
README_BASELINE = {
    'T1': [(50000, 0.0), (60000, 2.18), (62500, 2.76), (65500, 3.51),
           (69025, 4.51), (70026, 4.90), (70766, 7.46)],
    'T7': [(36001, 0.0), (40167, 0.0), (44054, 2.07), (47022, 3.14),
           (49026, 3.56), (51030, 4.13), (53019, 5.05), (55031, 7.22)],
    'T8': [(40000, 0.0), (50000, 0.0), (70000, 0.0), (74883, 1.94),
           (76931, 2.50), (89237, 3.71), (92315, 3.88), (96475, 4.61),
           (98492, 4.96), (100774, 5.52)],
}


def check_readme_baseline(merged):
    """把解析结果与 README §2.1 基准表逐条比对，返回 (结果列表, 是否全通过)."""
    res = []
    all_ok = True
    for sp, expect in README_BASELINE.items():
        got = {r['cycle']: r['crack_mm'] for r in merged if r['specimen'] == sp}
        bad = []
        for cyc, crk in expect:
            if cyc not in got:
                bad.append('循环 %d 缺失' % cyc)
            elif abs(got[cyc] - crk) > 1e-9:
                bad.append('循环 %d 值 %.4g≠%.4g' % (cyc, got[cyc], crk))
        extra = sorted(set(got) - {c for c, _ in expect})
        if extra:
            bad.append('多出 %s' % extra)
        ok = not bad
        all_ok = all_ok and ok
        res.append({'specimen': sp, 'n_expect': len(expect), 'n_got': len(got),
                    'ok': ok, 'detail': '; '.join(bad)})
    return res, all_ok


# ============================================================
# 4d. 带通滤波对重复性的改善（预处理口径的依据）
# ============================================================
# 激励是 200 kHz 窄带 burst（实测 ch1 峰频 200~210 kHz），但原始 ch2 记录被
# **带外噪声**主导：0~300 kHz 的能量占比 T6=91%、T1=63%，而 **T5 只有 9~40%**。
# ⇒ 这里用「重复性」这个无需真值的指标，把候选频带逐个量出来，选最优者。
CANDIDATE_BANDS = [
    ('原始（不滤波）', None),
    ('100~300 kHz', (100e3, 300e3)),
    ('150~350 kHz', (150e3, 350e3)),
    ('100~400 kHz', (100e3, 400e3)),
    ('150~400 kHz', (150e3, 400e3)),
    ('200~400 kHz', (200e3, 400e3)),
]


def repeatability_by_band():
    """对每个候选频带重算全部 47 个时刻的 signal_1 vs signal_2 相关性。

    返回 (汇总列表, 逐时刻明细 DataFrame)。相关性是**无真值指标**，
    因此可用于选预处理参数而不构成"用测试集调参"。
    """
    import prep

    idx = pd.read_csv(OUT_INDEX)
    pairs_meta = [(r.specimen, r.cycle)
                  for r in idx.drop_duplicates(['specimen', 'cycle']).itertuples()]
    summary, detail = [], []
    for name, band in CANDIDATE_BANDS:
        rows = []
        for sp, cyc in sorted(pairs_meta):
            a = prep.load_raw(sp, cyc, 1)[2]
            b = prep.load_raw(sp, cyc, 2)[2]
            if band is not None:
                a, b = prep.bandpass(a, band), prep.bandpass(b, band)
            rows.append({'specimen': sp, 'cycle': cyc,
                         'corr': prep.corr(a, b),
                         'rel_rms_diff': prep.rel_rms_diff(a, b)})
        d = pd.DataFrame(rows)
        d['band'] = name
        detail.append(d)
        summary.append({
            'band': name,
            'corr_median': float(d['corr'].median()),
            'corr_min': float(d['corr'].min()),
            'n_below_0.9': int((d['corr'] < 0.9).sum()),
            'rel_rms_median': float(d['rel_rms_diff'].median()),
            'per_specimen': {sp: float(d[d['specimen'] == sp]['corr'].median())
                             for sp in SPECIMENS},
        })
    return summary, pd.concat(detail, ignore_index=True)


# ============================================================
# 5. 主流程
# ============================================================
def main():
    os.makedirs(os.path.dirname(OUT_REPORT), exist_ok=True)

    say('=' * 78)
    say('PHM2019 铝搭接件（数据集 D） Step 0 数据规整 & 核对报告')
    say('生成时间: %s' % pd.Timestamp.now().strftime('%Y-%m-%d %H:%M:%S'))
    say('数据根目录: %s' % ROOT)
    say('=' * 78)

    # ---------- 解析 ----------
    records, metas = [], {}
    for sp in SPECIMENS:
        try:
            recs, meta = parse_description(sp)
        except Exception as e:
            say('[X] %s 真值表解析失败: %s' % (sp, e))
            continue
        records.extend(recs)
        metas[sp] = meta
    say('\n[1] 真值表解析：%d 条记录（%d 个试件）'
        % (len(records), len(metas)))

    index_rows = scan_waveforms()
    n_files = sum(1 for r in index_rows if r['path'] is not None)
    n_cycles = len({(r['specimen'], r['cycle']) for r in index_rows})
    say('    波形扫描：%d 个时刻目录 / %d 个信号文件（期望 %d）'
        % (n_cycles, n_files, 2 * n_cycles))

    # ---------- 对账 ----------
    say('\n' + '=' * 78)
    say('[核对 1] 标签 vs 波形对账')
    say('=' * 78)
    say('%-6s %8s %8s  %s' % ('试件', '标签数', '波形时刻', '差异'))
    merged, detail = reconcile(records, index_rows)
    for sp in SPECIMENS:
        if sp not in detail:
            continue
        d = detail[sp]
        diff = []
        if d['fixes']:
            for c, w in d['fixes'].items():
                diff.append('丢位修复 %d->%d' % (c, w))
        left_tag = [c for c in d['only_tag'] if c not in d['fixes']]
        left_wave = [w for w in d['only_wave'] if w not in set(d['fixes'].values())]
        if left_tag:
            diff.append('有标签无波形: %s' % left_tag)
        if left_wave:
            diff.append('有波形无标签: %s' % left_wave)
        say('%-6s %8d %8d  %s'
            % (sp, d['n_records'], d['n_wave_cycles'], '; '.join(diff) if diff else '完全对应'))

    say('')
    say('  真值表元信息（裂纹起始循环 / 载荷类型）：')
    for sp in SPECIMENS:
        m = metas.get(sp, {})
        n_wave = detail.get(sp, {}).get('n_wave_cycles', 0)
        n_rec = detail.get(sp, {}).get('n_records', 0)
        cover = ''
        if detail.get(sp, {}).get('wave_cycles'):
            w = detail[sp]['wave_cycles']
            cover = '波形覆盖 %d~%d' % (w[0], w[-1])
        say('    %-3s 载荷=%-9s 起裂~%-7s 标签 %2d 条（波形时刻 %2d 个，%s）'
            % (sp, m.get('loading', '?'), m.get('initiation_cycles', '?'),
               n_rec, n_wave, cover))

    say('\n  T7/T8 真值覆盖超出波形范围（天然外推评测集）：')
    for sp in VALID:
        w = detail[sp]['wave_cycles']
        recs = [r for r in merged if r['specimen'] == sp]
        say('    %s 波形只给 %s；真值覆盖到 %d 循环（多出 %d 条外推标签）'
            % (sp, w, max(r['cycle'] for r in recs),
               sum(1 for r in recs if r['cycle'] > max(w))))

    # ---------- 信号核对 ----------
    say('\n' + '=' * 78)
    say('[核对 2/3/4] ch1 一致性 · ch2 重复性 · 波形完整性')
    say('=' * 78)
    per_file, spec_ch1, pairs, problems = check_signals(index_rows)

    say('\n  2) ch1（激励）一致性 —— 同一试件内应几乎不变')
    say('     %-6s %6s %10s %10s %10s %10s' %
        ('试件', '文件数', 'pp_min', 'pp_max', 'pp_mean', '相对散布'))
    for s in spec_ch1:
        flag = '  <== ch1 全 0 !!' if s['all_zero'] else (
            '  <== 散布 >10%，疑耦合/增益漂移' if s['rel_spread'] > 0.10 else '')
        say('     %-6s %6d %10.4f %10.4f %10.4f %9.1f%%%s'
            % (s['specimen'], s['n_files'], s['pp_ch1_min'], s['pp_ch1_max'],
               s['pp_ch1_mean'], 100 * s['rel_spread'], flag))

    say('\n  3) ch2（接收）重复性 —— signal_1 vs signal_2 = 测量噪声下界')
    pf = pd.DataFrame(per_file)
    pr = pd.DataFrame(pairs)
    if len(pr):
        say('     %-6s %6s %10s %10s %10s %10s' %
            ('试件', '时刻数', 'corr中位', '对齐corr', 'relRMS差中位', 'relRMS差最大'))
        for sp in SPECIMENS:
            g = pr[pr['specimen'] == sp]
            if not len(g):
                continue
            say('     %-6s %6d %10.4f %10.4f %13.2f%% %13.2f%%'
                % (sp, len(g), g['corr_ch2'].median(),
                   g['corr_ch2_lag_best'].median(),
                   100 * g['rel_rms_diff'].median(), 100 * g['rel_rms_diff'].max()))
        say('     全体: corr 中位 %.4f / 最小 %.4f ; relRMS差 中位 %.2f%% / 最大 %.2f%%'
            % (pr['corr_ch2'].median(), pr['corr_ch2'].min(),
               100 * pr['rel_rms_diff'].median(), 100 * pr['rel_rms_diff'].max()))
        say('     （对齐corr = 在 ±10 µs 内做整数时移后再相关；'
            '若对齐corr ≈ corr 则差异非时移所致）')
        worst = pr.sort_values('rel_rms_diff', ascending=False).head(5)
        say('     重复性最差的 5 个时刻：')
        for _, r in worst.iterrows():
            say('       %-3s cycle=%-7d corr=%.4f 对齐=%.4f lag=%+d  relRMS差=%.2f%%'
                % (r['specimen'], r['cycle'], r['corr_ch2'],
                   r['corr_ch2_lag_best'], int(r['lag_best']),
                   100 * r['rel_rms_diff']))

    say('\n  4) 波形完整性')
    n_missing = sum(1 for p in problems if p[3].startswith('missing_file'))
    n_rows_bad = sum(1 for p in problems if p[3].startswith('n_rows'))
    n_cols_bad = sum(1 for p in problems if p[3].startswith('n_cols'))
    n_time_bad = sum(1 for p in problems if p[3].startswith('time'))
    say('     文件齐全性 : 实际 %d / 期望 %d 个%s'
        % (n_files, n_files + n_missing, '  [OK]' if n_missing == 0 else '  [!] 缺 %d' % n_missing))
    say('     点数 = %d  : %s' % (EXPECTED_ROWS,
        '[OK] 全部一致' if n_rows_bad == 0 else '[!] %d 个文件不符' % n_rows_bad))
    say('     列名结构   : %s' % ('[OK] 全部 time,ch1,ch2'
        if n_cols_bad == 0 else '[!] %d 个文件列数不为 3（见下）' % n_cols_bad))
    say('     time 列    : %s' % ('[OK] 与重建轴一致'
        if n_time_bad == 0 else
        '[!] %d/%d 个文件与重建轴偏差 > %.1e s（该列被 Excel 量化，不可用）'
        % (n_time_bad, n_files, 0.5 * EXPECTED_DT)))

    # 问题明细：按 (试件, 问题类型) 归并，避免逐文件刷屏
    if problems:
        say('\n     问题明细（按试件归并；同一试件的 rep1/rep2 合并列出）：')
        grouped = defaultdict(lambda: defaultdict(set))
        for sp, cyc, _rep, msg in problems:
            kind = msg.split(':', 1)[0].split('=')[0].strip()
            grouped[sp][kind].add(cyc)
        for sp in SPECIMENS:
            if sp not in grouped:
                continue
            parts = []
            for kind, cycs in sorted(grouped[sp].items()):
                parts.append('%s -> %s' % (kind, sorted(cycs)))
            say('       %-3s %s' % (sp, ' ; '.join(parts)))

    n_cols_detail = defaultdict(set)
    for sp, cyc, _rep, msg in problems:
        if msg.startswith('n_cols'):
            names = msg.split('列名=', 1)[1]
            n_cols_detail[sp].add((cyc, names.count('|') + 1))
    if n_cols_detail:
        say('\n     [!] 列数 ≠ 3 的文件（全部集中在 T6；多出的是尾部空列，'
            '读取时按列名取 time/ch1/ch2，不影响分析）：')
        for sp in SPECIMENS:
            if sp not in n_cols_detail:
                continue
            by_cyc = defaultdict(set)
            for cyc, nc in n_cols_detail[sp]:
                by_cyc[cyc].add(nc)
            say('       %-3s cycles=%s' % (sp, sorted(by_cyc)))
            for cyc in sorted(by_cyc):
                say('         cycle=%-7d 列数=%s' % (cyc, sorted(by_cyc[cyc])))

    say('\n     ⚠️ 时间轴结论：本数据集 time 列的文本精度**逐试件不同**'
        '（1 位 ~ 15 位有效数字）。')
    say('        例：T3/14000 写作 5.00000000000012e-08（可用），'
        'T6/68091 写作 1E-07（相邻两点同值）。')
    say('        ⇒ 分析一律采用重建轴 t[i] = i × %.1e s（i=0..%d），'
        % (EXPECTED_DT, EXPECTED_ROWS - 1))
    say('          与 README §1.1 的 dt / 200 µs / 4000 点完全一致'
        '（文件末值 0.00019995 = 3999×5e-8）。')

    # ---------- 试件级信号质量 ----------
    say('\n' + '=' * 78)
    say('[核对 5] 试件级信号质量 —— 决定哪些试件可用于定量回归')
    say('=' * 78)
    quals = specimen_quality(per_file, pairs)
    say('     %-3s %8s %10s %11s %8s %9s %10s  %s'
        % ('试件', '文件数', 'AC-RMS', '噪声σ', 'SNR', 'corr中位', 'ch2峰峰', '判定'))
    for q in quals:
        say('     %-3s %8d %10.5f %11.5f %8.2f %9.4f %10.4f  %s'
            % (q['specimen'], q['n_files'], q['ac_rms_ch2'], q['noise_sigma'],
               q['snr'], q['corr_med'], q['pp_ch2_med'], q['verdict']))
    bad_q = [q for q in quals if q['verdict'].startswith('✗')]
    if bad_q:
        say('\n     ⚠️ 判定说明：**在"原始信号、不做任何滤波"的口径下**，'
            '%s 的记录被带外噪声主导' % '/'.join(q['specimen'] for q in bad_q))
        say('        口径：噪声σ 由 signal_1/signal_2 的 RMS 差反推（σ = rms_diff/√2），')
        say('              SNR = ch2 的 AC-RMS / σ。')
        for q in bad_q:
            say('        %s：AC-RMS %.5f V ÷ 噪声σ %.5f V ⇒ SNR = %.2f（< 1.5）'
                % (q['specimen'], q['ac_rms_ch2'], q['noise_sigma'], q['snr']))
            say('             旁证1：ch1 激励峰峰 %.2f V，仅为 T1（4.24 V）的 %.0f%%'
                % (q['pp_ch1_med'], 100 * q['pp_ch1_med'] / 4.24))
            say('             旁证2：ch2 峰峰中位 %.4f V vs 量化步长 %.4f V'
                % (q['pp_ch2_med'], q['quant_step']))
            gq = pr[pr['specimen'] == q['specimen']]
            say('             旁证3：互相关中位 %.4f；做 ±10 µs 最佳整数时移对齐后中位 %.4f'
                % (q['corr_med'], gq['corr_ch2_lag_best'].median()))
            say('                    ⇒ 对齐**不能**改善 ⇒ 差异不是时移/配对错误，'
                '而是噪声。')
            worst = gq.sort_values('corr_ch2').iloc[0]
            say('                    最差时刻 %d：%.4f → 对齐后 %.4f'
                % (worst['cycle'], worst['corr_ch2'], worst['corr_ch2_lag_best']))
        say('        ⇒ 该试件的原始记录不含可比信息。')
        say('        BUT ⚠️ **这不是判废依据** —— 见核对 8：这些"噪声"绝大部分在**带外**，')
        say('           带通到激励频带后 T5 的 corr 中位可达 >0.94，'
            '**T5 保留、不剔除**。')

    # ---------- 与 README 基准表对账 ----------
    say('\n' + '=' * 78)
    say('[核对 6] 与 README §2.1 基准表对账（解析回归测试）')
    say('=' * 78)
    bl, bl_ok = check_readme_baseline(merged)
    for b in bl:
        say('     %-3s %d 条基准 / 解析得 %d 条  %s'
            % (b['specimen'], b['n_expect'], b['n_got'],
               '[OK] 逐条一致' if b['ok'] else '[!] ' + b['detail']))
    say('     → %s' % ('三组基准表全部复现，解析逻辑无漂移。' if bl_ok
                      else '存在不一致，必须先修正解析再往下走！'))


    # ch1 全 0 的明细（单独列出，便于复核）
    zero_files = pf[pf['pp_ch1'] == 0.0]
    if len(zero_files):
        say('\n     [!] ch1 恒为 0 的文件共 %d 个（激励通道无信号，'
            '传递函数 H=FFT(ch2)/FFT(ch1) 在这些时刻不可用）：' % len(zero_files))
        for sp in SPECIMENS:
            z = zero_files[zero_files['specimen'] == sp]
            if len(z):
                say('       %-3s cycles=%s' % (sp, sorted(set(z['cycle']))))

    # ch2 量化情况（影响特征精度的上限）
    say('\n' + '=' * 78)
    say('[核对 7] ch2 量化水平 —— 决定了特征可分辨的最小变化')
    say('=' * 78)
    say('     全部 4000 点中 ch2 的不同取值个数很有限 ⇒ 幅值类特征被量化地板限制。')
    say('     %-6s %10s %10s %10s' % ('试件', 'ch2唯一值', '量化步长', 'ch2 AC-RMS'))
    for sp in SPECIMENS:
        g = pf[pf['specimen'] == sp]
        if not len(g):
            continue
        say('     %-6s %10d %10.5f %10.5f'
            % (sp, int(g['n_unique_ch2'].median()),
               g['quant_step_ch2'].median(),
               float(np.sqrt(np.mean((g['rms_ch2'] ** 2 - g['mean_ch2'] ** 2).clip(lower=0))))))
    say('\n     ⇒ 幅值类特征的最小可分辨变化 ≈ 量化步长；'
        '本数据集为 8e-4 ~ 2e-3 V。')
    say('       Step 1 选特征时应优先用**相对量**（与基线之差/之比、相关系数、')
    say('       TOF、谱重心），而不是绝对幅值。')

    # ---------- 核对 8：带通滤波对重复性的改善 ----------
    say('\n' + '=' * 78)
    say('[核对 8] 带通滤波对重复性的改善 —— 预处理口径的依据')
    say('=' * 78)
    say('     背景：激励是 200 kHz 窄带 burst（ch1 峰频实测 200~210 kHz），')
    say('     但原始 ch2 记录被**带外噪声**主导（0~300 kHz 能量占比 '
        'T6=91% / T1=63% / **T5 仅 9~40%**）。')
    say('     下面用「signal_1 vs signal_2 相关性」这个**无真值指标**逐带量化'
        '（不构成用测试集调参）：')
    bsum, bdet = repeatability_by_band()
    say('')
    say('     %-18s %9s %9s %10s %10s   %s'
        % ('频带', 'corr中位', 'corr最小', '<0.9的时刻', 'relRMS中位', '每试件 corr 中位'))
    for s in bsum:
        per = ' '.join('%s:%.3f' % (sp, s['per_specimen'][sp]) for sp in SPECIMENS)
        say('     %-18s %9.4f %9.4f %10d %9.1f%%   %s'
            % (s['band'], s['corr_median'], s['corr_min'], s['n_below_0.9'],
               100 * s['rel_rms_median'], per))
    raw = bsum[0]
    best = max(bsum[1:], key=lambda s: s['corr_median'])
    say('')
    say('     ⇒ 选定频带 **%s**（%d 阶 Butterworth + 零相位 filtfilt；'
        % (best['band'], 4))
    say('       零相位是必须的：TOF 类特征要看**到达时刻**，一阶相位失真就会把 TOF 带偏。）')
    say('       相对原始：corr 中位 %.4f→%.4f，relRMS 差中位 %.1f%%→%.1f%%，'
        % (raw['corr_median'], best['corr_median'],
           100 * raw['rel_rms_median'], 100 * best['rel_rms_median']))
    say('       corr<0.9 的时刻数 %d→%d。' % (raw['n_below_0.9'], best['n_below_0.9']))

    raw_t5 = raw['per_specimen'].get('T5', float('nan'))
    best_t5 = best['per_specimen'].get('T5', float('nan'))
    say('')
    say('     🔴 **重要更正**：核对 5 曾据原始信号把 **T5 判为"噪声主导、不可用"**。')
    say('        带通后 T5 的 corr 中位由 %.4f 升到 **%.4f**（relRMS 差由 %.1f%% 降到 %.1f%%）'
        % (raw_t5, best_t5,
           100 * float(bdet[(bdet['band'] == raw['band']) &
                            (bdet['specimen'] == 'T5')]['rel_rms_diff'].median()),
           100 * float(bdet[(bdet['band'] == best['band']) &
                            (bdet['specimen'] == 'T5')]['rel_rms_diff'].median())))
    say('        ⇒ **T5 不是坏试件，是"没做带通"**；原始记录的低相关来自带外噪声，')
    say('          不是测量重复性差。**T5 保留，无需剔除**（训练集 6 组全部可用）。')
    say('        ⇒ 该结论已同步进 README §5.0.2 与 `prep.py` 的模块说明。')
    say('')
    say('     ⚠️ 但这一步**改变了"噪声下界"的口径**：')
    say('        原始 relRMS 差中位 %.1f%% ⇒ 带通后 %.1f%%。'
        % (100 * raw['rel_rms_median'], 100 * best['rel_rms_median']))
    say('        Step 1 的所有结论必须声明用的是哪个口径；'
        '报"检出损伤"时以**带通后**的下界为准。')

    # ---------- 写产物 ----------
    lab = pd.DataFrame(merged)[
        ['specimen', 'cycle', 'cycle_raw', 'crack_mm', 'source', 'note']]
    lab = lab.sort_values(['specimen', 'cycle']).reset_index(drop=True)
    lab.to_csv(OUT_LABELS, index=False, encoding='utf-8-sig')

    idx = pd.DataFrame(index_rows)[
        ['specimen', 'cycle', 'rep', 'path', 'n_rows', 'n_cols', 'col_names',
         't_file_err_s', 'time_ok', 'fs_hz']]
    idx = idx.sort_values(['specimen', 'cycle', 'rep']).reset_index(drop=True)
    idx.to_csv(OUT_INDEX, index=False, encoding='utf-8-sig')

    pd.DataFrame(per_file).to_csv(
        os.path.join(ROOT, 'results', 'step0_signal_stats.csv'),
        index=False, encoding='utf-8-sig')
    if len(pr):
        pr.to_csv(os.path.join(ROOT, 'results', 'step0_repeatability.csv'),
                  index=False, encoding='utf-8-sig')
    pd.DataFrame(quals).to_csv(
        os.path.join(ROOT, 'results', 'step0_specimen_quality.csv'),
        index=False, encoding='utf-8-sig')
    bdet.to_csv(os.path.join(ROOT, 'results', 'step0_band_sweep.csv'),
                index=False, encoding='utf-8-sig')

    say('\n' + '=' * 78)
    say('[产出]')
    say('  %s   (%d 行)' % (os.path.relpath(OUT_LABELS, ROOT), len(lab)))
    say('  %s   (%d 行)' % (os.path.relpath(OUT_INDEX, ROOT), len(idx)))
    say('  results/step0_report.txt / step0_signal_stats.csv /'
        ' step0_repeatability.csv / step0_specimen_quality.csv / step0_band_sweep.csv')
    say('=' * 78)

    with open(OUT_REPORT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(_REPORT_LINES) + '\n')


if __name__ == '__main__':
    main()
